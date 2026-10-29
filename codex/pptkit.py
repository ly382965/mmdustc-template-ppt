"""Thin OOXML layer over python-pptx.

Everything is written as explicit DrawingML (no theme style references), so a shape looks the same
in every viewer: text runs carry both a Latin and an East-Asian face, shapes carry explicit fill and
line, SVG art is embedded as svgBlip with a PNG fallback.
"""
import copy
import hashlib
import io
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu

from . import tokens as T
from .svg import rasterize, svg_size

C_NS = 'http://schemas.openxmlformats.org/drawingml/2006/chart'

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}
NSDECL = ' '.join('xmlns:%s="%s"' % kv for kv in NS.items())
SVG_EXT = '{96DAC541-7B7A-43D3-8B79-37D633B846F1}'
ASVG = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'


def E(inches):
    return int(round(inches * 914400))


def xml(s):
    return etree.fromstring(s)


# ======================================================================================
#  Text model
# ======================================================================================
class R:
    """A text run. `role` picks the (Latin, East-Asian) faces from tokens.FONTS."""

    def __init__(self, text, role='body', size=None, color=None, bold=False, italic=False,
                 spc=None, underline=False, baseline=None, font=None, field=None):
        self.text, self.role, self.size, self.color = text, role, size, color
        self.bold, self.italic, self.spc, self.underline = bold, italic, spc, underline
        self.baseline, self.font, self.field = baseline, font, field


class P:
    """A paragraph: runs + layout. `line` is a multiple (1.0 = single); `before`/`after` in pt."""

    def __init__(self, runs, align='l', line=None, before=0, after=0, indent=0.0, hang=0.0,
                 bullet=None, bullet_color=None, bullet_font=None, bullet_size=None, size=None,
                 role=None, color=None, bold=None, spc=None, first=0.0, exact=None, right=0.0):
        if isinstance(runs, (str, R)):
            runs = [runs]
        self.runs = [r if isinstance(r, R) else R(r) for r in runs]
        for r in self.runs:  # paragraph-level defaults fill unset run attributes
            if size is not None and r.size is None:
                r.size = size
            if role is not None and r.role == 'body':
                r.role = role
            if color is not None and r.color is None:
                r.color = color
            if bold is not None and r.bold is False:
                r.bold = bold
            if spc is not None and r.spc is None:
                r.spc = spc
        self.align, self.line, self.before, self.after = align, line, before, after
        self.indent, self.hang, self.first, self.exact, self.right = indent, hang, first, exact, right
        self.bullet, self.bullet_color = bullet, bullet_color
        self.bullet_font, self.bullet_size = bullet_font, bullet_size


def rpr_xml(r, tag='a:rPr'):
    latin, ea = (r.font, r.font) if r.font else T.FONTS[r.role]
    a = ['lang="zh-CN"', 'altLang="en-US"', 'sz="%d"' % int(round((r.size or 16) * 100)),
         'b="%d"' % bool(r.bold), 'i="%d"' % bool(r.italic), 'dirty="0"']
    spc = r.spc if r.spc is not None else (T.TRACK.get(r.role, 0) * (r.size or 16) * 100 if not r.font else 0)
    if spc:
        a.append('spc="%d"' % int(round(spc)))
    if r.underline:
        a.append('u="sng"')
    if r.baseline:
        a.append('baseline="%d"' % int(r.baseline))
    fill = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % (r.color or T.TEXT)
    faces = ('<a:latin typeface="%s"/><a:ea typeface="%s"/><a:cs typeface="%s"/>'
             % (escape(latin), escape(ea), escape(latin)))
    return '<%s %s>%s%s</%s>' % (tag, ' '.join(a), fill, faces, tag)


def p_xml(p):
    algn = {'l': 'l', 'c': 'ctr', 'r': 'r', 'j': 'just', 'd': 'dist'}[p.align]
    attrs = ['algn="%s"' % algn, 'eaLnBrk="1" latinLnBrk="0" hangingPunct="1"']
    if p.indent or p.hang:
        attrs.append('marL="%d" indent="%d"' % (E(p.indent + p.hang), -E(p.hang)))
    elif p.first:
        attrs.append('indent="%d"' % E(p.first))
    if p.right:
        attrs.append('marR="%d"' % E(p.right))
    kids = []
    if p.exact:
        kids.append('<a:lnSpc><a:spcPts val="%d"/></a:lnSpc>' % int(p.exact * 100))
    elif p.line:
        kids.append('<a:lnSpc><a:spcPct val="%d"/></a:lnSpc>' % int(p.line * 100000))
    kids.append('<a:spcBef><a:spcPts val="%d"/></a:spcBef>' % int(p.before * 100))
    kids.append('<a:spcAft><a:spcPts val="%d"/></a:spcAft>' % int(p.after * 100))
    if p.bullet:
        kids.append('<a:buClr><a:srgbClr val="%s"/></a:buClr>' % (p.bullet_color or T.RUBRIC))
        kids.append('<a:buSzPct val="%d"/>' % int((p.bullet_size or 0.8) * 100000))
        kids.append('<a:buFont typeface="%s"/>' % escape(p.bullet_font or 'Segoe UI Symbol'))
        kids.append('<a:buChar char="%s"/>' % escape(p.bullet))
    else:
        kids.append('<a:buNone/>')
    runs = []
    for i, r in enumerate(p.runs):
        if r.field:
            runs.append('<a:fld id="{5C0F3E84-4E8B-4E57-9D42-3A4C2E1B7A%02d}" type="%s">%s<a:t>%s</a:t></a:fld>'
                        % (i, r.field, rpr_xml(r), escape(r.text)))
            continue
        for j, chunk in enumerate(r.text.split('\n')):
            if j:
                runs.append('<a:br>%s</a:br>' % rpr_xml(r))
            if chunk:
                runs.append('<a:r>%s<a:t>%s</a:t></a:r>' % (rpr_xml(r), escape(chunk)))
    runs = ''.join(runs)
    end = rpr_xml(p.runs[-1], 'a:endParaRPr') if p.runs else ''
    return '<a:p><a:pPr %s>%s</a:pPr>%s%s</a:p>' % (' '.join(attrs), ''.join(kids), runs, end)


def body_xml(paras, anchor='t', inset=(0, 0, 0, 0), autofit=None, vert=None, wrap=True, tag='p:txBody'):
    l, t, r, b = inset
    fit = {'shape': '<a:spAutoFit/>', 'shrink': '<a:normAutofit/>'}.get(autofit, '<a:noAutofit/>')
    va = ' vert="%s"' % vert if vert else ''
    return ('<%s><a:bodyPr wrap="%s" lIns="%d" tIns="%d" rIns="%d" bIns="%d" anchor="%s" rtlCol="0"%s>%s</a:bodyPr>'
            '<a:lstStyle/>%s</%s>' % (tag, 'square' if wrap else 'none', E(l), E(t), E(r), E(b),
                                      {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor], va, fit,
                                      ''.join(p_xml(p) for p in paras), tag))

# ======================================================================================
#  Fills and lines
# ======================================================================================
def fill_xml(fill, alpha=None):
    if not fill:
        return '<a:noFill/>'
    a = '<a:alpha val="%d"/>' % int(alpha * 100000) if alpha is not None else ''
    return '<a:solidFill><a:srgbClr val="%s">%s</a:srgbClr></a:solidFill>' % (fill, a)


def ln_xml(color=None, w=0.75, dash=None, head=None, tail=None, cap='rnd'):
    if not color:
        return '<a:ln><a:noFill/></a:ln>'
    d = '<a:prstDash val="%s"/>' % dash if dash else ''
    he = '<a:headEnd type="%s" w="sm" len="sm"/>' % head if head else ''
    te = '<a:tailEnd type="%s" w="sm" len="sm"/>' % tail if tail else ''
    return ('<a:ln w="%d" cap="%s"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>%s<a:round/>%s%s</a:ln>'
            % (int(w * 12700), cap, color, d, he, te))


# ======================================================================================
#  Slide wrapper
# ======================================================================================
class Canvas:
    """Drawing surface over a slide, layout or master part."""

    def __init__(self, deck, part, sptree):
        self.deck, self.part, self.tree = deck, part, sptree
        ids = [int(e.get('id')) for e in sptree.iter('{%s}cNvPr' % NS['p'])]
        self._id = max(ids + [1])

    def _nid(self):
        self._id += 1
        return self._id

    def _append(self, s):
        el = xml(s)
        self.tree.append(el)
        return el

    @staticmethod
    def _xfrm(x, y, w, h, rot=0, flipH=False, flipV=False, tag='a:xfrm'):
        a = ''
        if rot:
            a += ' rot="%d"' % int(rot * 60000)
        if flipH:
            a += ' flipH="1"'
        if flipV:
            a += ' flipV="1"'
        return '<%s%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></%s>' % (tag, a, E(x), E(y), E(w), E(h), tag)

    # -- text ---------------------------------------------------------------------------
    def text(self, x, y, w, h, paras, name='Text', anchor='t', inset=(0, 0, 0, 0), autofit=None,
             fill=None, line=None, lw=0.75, geom='rect', radius=None, vert=None, wrap=True, alpha=None):
        if isinstance(paras, (P, R, str)):
            paras = [paras if isinstance(paras, P) else P(paras)]
        av = '<a:gd name="adj" fmla="val %d"/>' % int(radius * 100000) if radius is not None else ''
        s = ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
             '<p:spPr>%s<a:prstGeom prst="%s"><a:avLst>%s</a:avLst></a:prstGeom>%s%s</p:spPr>%s</p:sp>'
             % (NSDECL, self._nid(), escape(name), self._xfrm(x, y, w, h), geom, av, fill_xml(fill, alpha),
                ln_xml(line, lw), body_xml(paras, anchor, inset, autofit, vert, wrap)))
        return self._append(s)

    # -- shapes -------------------------------------------------------------------------
    def shape(self, geom, x, y, w, h, fill=None, line=None, lw=0.75, dash=None, name='Shape', alpha=None,
              radius=None, rot=0, adj=None, paras=None, anchor='m', inset=(0.05, 0.03, 0.05, 0.03)):
        gd = ''
        if radius is not None:
            gd = '<a:gd name="adj" fmla="val %d"/>' % int(radius * 100000)
        if adj:
            gd = ''.join('<a:gd name="%s" fmla="val %d"/>' % kv for kv in adj.items())
        body = body_xml(paras, anchor, inset) if paras else ''
        s = ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
             '<p:spPr>%s<a:prstGeom prst="%s"><a:avLst>%s</a:avLst></a:prstGeom>%s%s</p:spPr>%s</p:sp>'
             % (NSDECL, self._nid(), escape(name), self._xfrm(x, y, w, h, rot), geom, gd,
                fill_xml(fill, alpha), ln_xml(line, lw, dash), body))
        return self._append(s)

    def rect(self, x, y, w, h, **kw):
        return self.shape('rect', x, y, w, h, **kw)

    def rrect(self, x, y, w, h, r=0.06, **kw):
        return self.shape('roundRect', x, y, w, h, radius=r, **kw)

    def oval(self, x, y, w, h, **kw):
        return self.shape('ellipse', x, y, w, h, **kw)

    def line(self, x1, y1, x2, y2, color=None, lw=0.75, dash=None, head=None, tail=None, name='Line', cap='rnd'):
        x, y = min(x1, x2), min(y1, y2)
        w, h = abs(x2 - x1), abs(y2 - y1)
        s = ('<p:cxnSp %s><p:nvCxnSpPr><p:cNvPr id="%d" name="%s"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
             '<p:spPr>%s<a:prstGeom prst="line"><a:avLst/></a:prstGeom>%s</p:spPr></p:cxnSp>'
             % (NSDECL, self._nid(), escape(name), self._xfrm(x, y, w, h, flipH=x2 < x1, flipV=y2 < y1),
                ln_xml(color or T.RULE, lw, dash, head, tail, cap)))
        return self._append(s)

    # -- pictures -----------------------------------------------------------------------
    def _pic(self, rid, x, y, w, h, name, svg_rid=None, crop=None, alpha=None):
        ext = ''
        if svg_rid:
            ext = ('<a:extLst><a:ext uri="%s"><asvg:svgBlip xmlns:asvg="%s" r:embed="%s"/></a:ext></a:extLst>'
                   % (SVG_EXT, ASVG, svg_rid))
        amod = '<a:alphaModFix amt="%d"/>' % int(alpha * 100000) if alpha is not None else ''
        src = ''
        if crop:
            l, t, r, b = crop
            src = '<a:srcRect l="%d" t="%d" r="%d" b="%d"/>' % tuple(int(v * 100000) for v in (l, t, r, b))
        s = ('<p:pic %s><p:nvPicPr><p:cNvPr id="%d" name="%s"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/>'
             '</p:cNvPicPr><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="%s">%s%s</a:blip>%s'
             '<a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr>%s<a:prstGeom prst="rect"><a:avLst/>'
             '</a:prstGeom></p:spPr></p:pic>'
             % (NSDECL, self._nid(), escape(name), rid, amod, ext, src, self._xfrm(x, y, w, h)))
        return self._append(s)

    def image(self, path_or_bytes, x, y, w=None, h=None, name='Picture', crop=None, alpha=None, fit=None):
        """Picture at (x, y). With both w and h: fit='contain' letterboxes inside the box (centred),
        fit='cover' crops the source to fill it; without fit the box size is used as given."""
        from PIL import Image
        blob = path_or_bytes if isinstance(path_or_bytes, bytes) else open(path_or_bytes, 'rb').read()
        _, rid = self.part.get_or_add_image_part(io.BytesIO(blob))
        with Image.open(io.BytesIO(blob)) as im:
            iw, ih = im.size
        if crop:
            iw, ih = iw * (1 - crop[0] - crop[2]), ih * (1 - crop[1] - crop[3])
        if fit and w and h:
            if fit == 'contain':
                s = min(w / iw, h / ih)
                x, y, w, h = x + (w - iw * s) / 2, y + (h - ih * s) / 2, iw * s, ih * s
            elif fit == 'cover':
                a_img, a_box = iw / ih, w / h
                e = (1 - a_box / a_img) / 2 if a_img > a_box else (1 - a_img / a_box) / 2
                crop = (e, 0, e, 0) if a_img > a_box else (0, e, 0, e)
        else:
            w, h = _fit(w, h, iw / ih)
        return self._pic(rid, x, y, w, h, name, crop=crop, alpha=alpha)

    def svg(self, svg, x, y, w=None, h=None, name='Ornament', dpi=300, alpha=None):
        sw, sh = svg_size(svg)
        w, h = _fit(w, h, sw / sh)
        png = self.deck.fallback_png(svg, min(3000, max(64, int(w * dpi))))
        _, rid = self.part.get_or_add_image_part(io.BytesIO(png))
        svg_rid = self.part.relate_to(self.deck.svg_part(svg), RT.IMAGE)
        return self._pic(rid, x, y, w, h, name, svg_rid=svg_rid, alpha=alpha)

    # -- tables -------------------------------------------------------------------------
    def table(self, x, y, col_w, row_h, cells, name='Table', borders=None, fills=None, margins=(0.08, 0.04),
              anchor='m'):
        """cells[r][c] = list of P; borders[(r, side)] = (color, pt) applied to every cell of row r;
        side in 'T'/'B'. fills[r] = hex or None."""
        borders, fills = borders or {}, fills or {}
        grid = ''.join('<a:gridCol w="%d"/>' % E(w) for w in col_w)
        rows = []
        for ri, row in enumerate(cells):
            tcs = []
            for ci, paras in enumerate(row):
                def edge(side):
                    spec = borders.get((ri, side))
                    if not spec:
                        return '<a:ln%s w="0"><a:noFill/></a:ln%s>' % (side, side)
                    c, pt = spec
                    return ('<a:ln%s w="%d" cap="flat" cmpd="sng"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
                            '<a:prstDash val="solid"/></a:ln%s>' % (side, int(pt * 12700), c, side))
                f = fills.get(ri)
                tc_fill = fill_xml(f) if f else '<a:noFill/>'
                mx, my = margins
                tcs.append('<a:tc>%s<a:tcPr marL="%d" marR="%d" marT="%d" marB="%d" anchor="%s">%s%s%s%s%s</a:tcPr></a:tc>'
                           % (body_xml(paras, tag='a:txBody'), E(mx), E(mx), E(my), E(my),
                              {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor],
                              edge('L'), edge('R'), edge('T'), edge('B'), tc_fill))
            rows.append('<a:tr h="%d">%s</a:tr>' % (E(row_h[ri]), ''.join(tcs)))
        s = ('<p:graphicFrame %s><p:nvGraphicFramePr><p:cNvPr id="%d" name="%s"/><p:cNvGraphicFramePr>'
             '<a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>%s'
             '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
             '<a:tbl><a:tblPr firstRow="1" bandRow="0"/><a:tblGrid>%s</a:tblGrid>%s</a:tbl></a:graphicData>'
             '</a:graphic></p:graphicFrame>'
             % (NSDECL, self._nid(), escape(name), self._xfrm(x, y, sum(col_w), sum(row_h), tag='p:xfrm'),
                grid, ''.join(rows)))
        return self._append(s)

    # -- charts -------------------------------------------------------------------------
    def line_chart(self, x, y, w, h, categories, series, y_min=None, y_max=None, y_step=None,
                   x_title=None, y_title=None, label=None, name='Chart', number_format='0'):
        """Native, editable line chart in template styling. series = [dict(name, values, color,
        width=1.75, dash=False, hollow=False)]; `label` = index of the series that gets value labels."""
        from pptx.chart.data import CategoryChartData
        from pptx.chart.datalabel import DataLabels
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
        from pptx.enum.chart import XL_TICK_MARK
        from pptx.enum.dml import MSO_LINE
        from pptx.util import Pt
        rgb = RGBColor.from_string
        cd = CategoryChartData()
        cd.categories = categories
        for s in series:
            cd.add_series(s['name'], s['values'])
        gf = self.slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, E(x), E(y), E(w), E(h), cd)
        gf.name = name
        self._id = max(self._id, gf.shape_id)
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.font.size = Pt(12)
        ch.font.color.rgb = rgb(T.MUTED)
        va, ca = ch.value_axis, ch.category_axis
        if y_min is not None:
            va.minimum_scale = y_min
        if y_max is not None:
            va.maximum_scale = y_max
        if y_step:
            va.major_unit = y_step
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(T.RULE)
        va.major_gridlines.format.line.width = Pt(0.6)
        va.format.line.fill.background()
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        ca.format.line.color.rgb = rgb(T.FAINT)
        ca.format.line.width = Pt(0.75)
        ca.major_tick_mark = XL_TICK_MARK.NONE
        for axis, title in ((va, y_title), (ca, x_title)):
            if title:
                axis.has_title = True
                tf = axis.axis_title.text_frame
                tf.text = title
                f = tf.paragraphs[0].runs[0].font
                f.size, f.bold, f.color.rgb = Pt(11.5), False, rgb(T.MUTED)
        for k, s in enumerate(ch.plots[0].series):
            spec, col = series[k], rgb(series[k]['color'])
            s.smooth = False
            s.format.line.color.rgb = col
            s.format.line.width = Pt(spec.get('width', 1.75))
            if spec.get('dash'):
                s.format.line.dash_style = MSO_LINE.DASH
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = spec.get('marker', 7)
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb('FBF8F1') if spec.get('hollow') else col
            s.marker.format.line.color.rgb = col
            s.marker.format.line.width = Pt(1.25)
            if label == k:
                dl = DataLabels(s._element.get_or_add_dLbls())
                dl.show_value = True
                dl.position = XL_LABEL_POSITION.ABOVE
                dl.number_format, dl.number_format_is_linked = '0.0', False
                dl.font.size, dl.font.color.rgb = Pt(11.5), col
        cs = ch._chartSpace
        c = '{%s}' % C_NS
        nofill = '<c:spPr xmlns:c="%s" %s><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>' % (C_NS, NSDECL)
        cs.find(c + 'chart').addnext(xml(nofill))
        cs.find(c + 'chart').find(c + 'plotArea').append(xml(nofill))
        latin, ea = T.FONTS['body']
        for tag in ('defRPr', 'rPr'):
            for rp in cs.iter('{%s}%s' % (NS['a'], tag)):
                for old in rp.findall('{%s}latin' % NS['a']) + rp.findall('{%s}ea' % NS['a']):
                    rp.remove(old)
                rp.append(xml('<a:latin %s typeface="%s"/>' % (NSDECL, latin)))
                rp.append(xml('<a:ea %s typeface="%s"/>' % (NSDECL, ea)))
        return gf


def _fit(w, h, aspect):
    if w is None and h is None:
        raise ValueError('need w or h')
    if w is None:
        w = h * aspect
    if h is None:
        h = w / aspect
    return w, h


# ======================================================================================
#  Deck
# ======================================================================================
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(E(T.W)), Emu(E(T.H))
        self._svg_parts, self._png_cache, self._n_svg = {}, {}, 0
        self.layouts = {}
        # keep a single blank layout as the prototype for our own layouts
        keep = self.prs.slide_layouts[6]
        for lay in list(self.prs.slide_layouts):
            if lay is not keep:
                self.prs.slide_layouts.remove(lay)
        self._proto = keep
        master = self.prs.slide_master
        for ph in list(master.placeholders):
            ph._element.getparent().remove(ph._element)

    def set_theme(self, name='mmdustc Codex'):
        """Theme colours = template palette, theme fonts = template serif (new shapes inherit them)."""
        part = self.prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        root.set('name', name)
        te = root.find('a:themeElements', NS)
        acc = ''.join('<a:accent%d><a:srgbClr val="%s"/></a:accent%d>' % (i + 1, c, i + 1)
                      for i, c in enumerate(T.SERIES[:6]))
        clr = xml('<a:clrScheme %s name="%s"><a:dk1><a:srgbClr val="%s"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/>'
                  '</a:lt1><a:dk2><a:srgbClr val="%s"/></a:dk2><a:lt2><a:srgbClr val="%s"/></a:lt2>%s'
                  '<a:hlink><a:srgbClr val="%s"/></a:hlink><a:folHlink><a:srgbClr val="%s"/></a:folHlink>'
                  '</a:clrScheme>' % (NSDECL, name, T.INK, T.TEXT, T.PAPER, acc, T.RUBRIC, T.SERIES[4]))
        te.replace(te.find('a:clrScheme', NS), clr)
        (mj_l, mj_e), (mn_l, mn_e) = T.FONTS['title'], T.FONTS['body']
        fonts = xml('<a:fontScheme %s name="%s"><a:majorFont><a:latin typeface="%s"/><a:ea typeface="%s"/>'
                    '<a:cs typeface=""/><a:font script="Hans" typeface="%s"/></a:majorFont><a:minorFont>'
                    '<a:latin typeface="%s"/><a:ea typeface="%s"/><a:cs typeface=""/>'
                    '<a:font script="Hans" typeface="%s"/></a:minorFont></a:fontScheme>'
                    % (NSDECL, name, mj_l, mj_e, mj_e, mn_l, mn_e, mn_e))
        te.replace(te.find('a:fontScheme', NS), fonts)
        part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

    # -- packaging helpers --------------------------------------------------------------------
    def svg_part(self, svg):
        key = hashlib.sha1(svg.encode('utf-8')).hexdigest()
        part = self._svg_parts.get(key)
        if part is None:
            self._n_svg += 1
            part = Part(PackURI('/ppt/media/art%d.svg' % self._n_svg), 'image/svg+xml',
                        self.prs.part.package, svg.encode('utf-8'))
            self._svg_parts[key] = part
        return part

    def fallback_png(self, svg, width_px):
        key = (hashlib.sha1(svg.encode('utf-8')).hexdigest(), width_px)
        if key not in self._png_cache:
            self._png_cache[key] = rasterize(svg, width_px)
        return self._png_cache[key]

    def set_background(self, part_owner, image_bytes):
        """Picture background on a master/layout/slide (cSld/bg)."""
        part = part_owner.part
        _, rid = part.get_or_add_image_part(io.BytesIO(image_bytes))
        cSld = part_owner._element.find('{%s}cSld' % NS['p'])
        old = cSld.find('{%s}bg' % NS['p'])
        if old is not None:
            cSld.remove(old)
        bg = xml('<p:bg %s><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="%s"/><a:srcRect/>'
                 '<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>' % (NSDECL, rid))
        cSld.insert(0, bg)

    def add_layout(self, name, draw=None):
        """Clone the blank prototype layout, rename it and let `draw(canvas)` decorate it."""
        src = self._proto
        if not self.layouts:
            lay = src
        else:
            lay = self._clone_layout(src)
        lay._element.find('{%s}cSld' % NS['p']).set('name', name)
        for ph in list(lay.placeholders):
            ph._element.getparent().remove(ph._element)
        if draw:
            draw(Canvas(self, lay.part, lay.shapes._spTree))
        self.layouts[name] = lay
        return lay

    def _clone_layout(self, src):
        from pptx.parts.slide import SlideLayoutPart
        from pptx.slide import SlideLayout
        pkg = self.prs.part.package
        master_part = self.prs.slide_master.part
        n = len(self.prs.slide_layouts) + 1
        while any(str(p.partname) == '/ppt/slideLayouts/slideLayout%d.xml' % n for p in pkg.iter_parts()):
            n += 1
        el = copy.deepcopy(src._element)
        for t in list(el.find('{%s}cSld' % NS['p']).find('{%s}spTree' % NS['p'])):
            if t.tag.endswith('}sp') or t.tag.endswith('}pic') or t.tag.endswith('}cxnSp'):
                t.getparent().remove(t)
        part = SlideLayoutPart(PackURI('/ppt/slideLayouts/slideLayout%d.xml' % n),
                               src.part.content_type, pkg, el)
        part.relate_to(master_part, RT.SLIDE_MASTER)
        rid = master_part.relate_to(part, RT.SLIDE_LAYOUT)
        lst = self.prs.slide_master._element.find('{%s}sldLayoutIdLst' % NS['p'])
        used = [int(e.get('id')) for e in lst] + [int(e.get('id')) for e in
                                                   self.prs._element.find('{%s}sldMasterIdLst' % NS['p'])]
        lst.append(xml('<p:sldLayoutId %s id="%d" r:id="%s"/>' % (NSDECL, max(used) + 1, rid)))
        return SlideLayout(el, part)

    def slide(self, layout):
        s = self.prs.slides.add_slide(self.layouts[layout])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        c = Canvas(self, s.part, s.shapes._spTree)
        c.slide = s
        return c

    def save(self, path):
        self.prs.save(path)
